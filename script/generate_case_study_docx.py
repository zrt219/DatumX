from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIR = REPO_ROOT / "output" / "doc"
LOGO_PATH = REPO_ROOT / "New folder" / "DatumX-removebg-preview.png"

PRIMARY = RGBColor(120, 163, 123)  # #78A37B
SECONDARY = RGBColor(95, 127, 167)  # #5F7FA7
TEXT = RGBColor(26, 33, 43)
MUTED = RGBColor(96, 108, 122)
RULE = RGBColor(198, 209, 198)
CALLOUT_FILL = "EEF4EE"
CODE_FILL = "F5F7F8"
TABLE_HEADER_FILL = "E6EEE7"

MOJIBAKE_REPLACEMENTS = {
    "\ufeff": "",
    "\u00a0": " ",
    "â€”": "-",
    "â€“": "-",
    "â€™": "'",
    "â€œ": '"',
    "â€": '"',
    "â€¦": "...",
    "â†’": "->",
    "Â·": "|",
    "Ã—": "x",
}


@dataclass(frozen=True)
class DocumentConfig:
    key: str
    source_path: Path
    output_path: Path
    doc_title: str
    cover_title: str
    cover_subtitle: str
    value_prop: str
    stack_line: str
    live_line: str
    static_contents_title: str
    intro_note: str
    callouts: dict[str, tuple[str, str]]


DOCS = [
    DocumentConfig(
        key="case-study",
        source_path=REPO_ROOT / "docs" / "datumx-case-study.md",
        output_path=OUTPUT_DIR / "DatumX-Case-Study.docx",
        doc_title="DatumX Protocol Case Study",
        cover_title="DatumX - Verifying AI-Transformed Exploration Data with On-Chain Provenance and Consensus Validation",
        cover_subtitle="Long-form protocol case study",
        value_prop=(
            "DatumX treats transformed industrial datasets as governed analytical releases with canonical payloads, "
            "deterministic lineage, validator consensus, and auditor-controlled finalization."
        ),
        stack_line="XRPL EVM Testnet | Solidity 0.8.24 | Foundry | React | Vite | TypeScript | viem",
        live_line="Live console: https://datumx.vercel.app | Canonical ValidationModule: 0xD60BE5B6A28eAF13De6DB6e02726551304E82baf",
        static_contents_title="Case Study Contents",
        intro_note="Deep protocol narrative for engineers, protocol designers, technical reviewers, and investment diligence.",
        callouts={
            "1. Executive Summary": (
                "Thesis",
                "DatumX is not trying to store geology on-chain. It is defining a release-governance boundary for transformed industrial data.",
            ),
            "6. Adapter Model (Key Innovation)": (
                "Typed Boundary",
                "The adapter boundary is where local project mess becomes canonical protocol input. V5 made that boundary explicit with AdapterPayload.",
            ),
            "8. Validation + Consensus System": (
                "Role Separation",
                "Validator-only reviews and auditor-only finalization are distinct on purpose. Consensus and closure are not the same action.",
            ),
            "11. XRPL EVM Deployment Model": (
                "Live Proof",
                "DatumX has already proven root-to-child lineage on XRPL EVM Testnet and restored thresholds after each fast-path exercise.",
            ),
        },
    ),
    DocumentConfig(
        key="executive-brief",
        source_path=REPO_ROOT / "docs" / "datumx-executive-technical-brief.md",
        output_path=OUTPUT_DIR / "DatumX-Executive-Technical-Brief.docx",
        doc_title="DatumX Executive Technical Brief",
        cover_title="DatumX - Executive Technical Brief",
        cover_subtitle="Leadership-facing protocol and deployment briefing",
        value_prop=(
            "DatumX converts AI-transformed exploration data from ambiguous files into governed releases with deterministic lineage, "
            "validator accountability, and on-chain finalization."
        ),
        stack_line="XRPL EVM Testnet | Solidity 0.8.24 | Foundry | React | Vite | TypeScript | viem | OpenZeppelin",
        live_line="Live console: https://datumx.vercel.app | Canonical V5 deploy manifest: deployments/xrpl-testnet.json",
        static_contents_title="Brief Contents",
        intro_note="Executive-technical summary calibrated for technical C-suite, strategic partners, and Web3 investment review.",
        callouts={
            "Executive Summary": (
                "Operating Claim",
                "DatumX is valuable because it governs transformed releases, not because it stores raw data more expensively.",
            ),
            "Chain and Infrastructure Specifics": (
                "XRPL Reality",
                "The deployment posture is deliberately XRPL-aware: legacy tx mode, explicit gas posture, and explorer-state honesty.",
            ),
            "Engineering Principles": (
                "Design Standard",
                "The system favors explicit invariants, custom errors, and deterministic progression over flexible but ambiguous workflow logic.",
            ),
        },
    ),
]


def normalize_text(text: str) -> str:
    for old, new in MOJIBAKE_REPLACEMENTS.items():
        text = text.replace(old, new)
    return text.rstrip()


def strip_markdown_inline(text: str) -> str:
    text = normalize_text(text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"\*([^*]+)\*", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1", text)
    return text.strip()


def set_section_margins(section) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)


def ensure_style(document: Document, name: str, style_type: WD_STYLE_TYPE):
    try:
        return document.styles[name]
    except KeyError:
        return document.styles.add_style(name, style_type)


def configure_styles(document: Document) -> None:
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.18

    body = ensure_style(document, "DatumX Body", WD_STYLE_TYPE.PARAGRAPH)
    body.base_style = normal
    body.font.name = "Arial"
    body.font.size = Pt(10.5)
    body.font.color.rgb = TEXT
    body.paragraph_format.space_after = Pt(7)
    body.paragraph_format.line_spacing = 1.18

    label = ensure_style(document, "DatumX Label", WD_STYLE_TYPE.PARAGRAPH)
    label.base_style = normal
    label.font.name = "Courier New"
    label.font.size = Pt(8.5)
    label.font.bold = True
    label.font.color.rgb = PRIMARY
    label.paragraph_format.space_after = Pt(3)

    section_title = ensure_style(document, "DatumX Section Title", WD_STYLE_TYPE.PARAGRAPH)
    section_title.base_style = normal
    section_title.font.name = "Courier New"
    section_title.font.size = Pt(14)
    section_title.font.bold = True
    section_title.font.color.rgb = TEXT
    section_title.paragraph_format.space_before = Pt(6)
    section_title.paragraph_format.space_after = Pt(6)
    section_title.paragraph_format.keep_with_next = True

    subsection = ensure_style(document, "DatumX Subsection", WD_STYLE_TYPE.PARAGRAPH)
    subsection.base_style = normal
    subsection.font.name = "Courier New"
    subsection.font.size = Pt(11)
    subsection.font.bold = True
    subsection.font.color.rgb = SECONDARY
    subsection.paragraph_format.space_before = Pt(4)
    subsection.paragraph_format.space_after = Pt(4)
    subsection.paragraph_format.keep_with_next = True

    meta = ensure_style(document, "DatumX Meta", WD_STYLE_TYPE.PARAGRAPH)
    meta.base_style = normal
    meta.font.name = "Courier New"
    meta.font.size = Pt(9)
    meta.font.color.rgb = MUTED
    meta.paragraph_format.space_after = Pt(4)

    callout = ensure_style(document, "DatumX Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.base_style = normal
    callout.font.name = "Arial"
    callout.font.size = Pt(10.5)
    callout.font.italic = True
    callout.font.color.rgb = TEXT
    callout.paragraph_format.space_after = Pt(4)
    callout.paragraph_format.line_spacing = 1.15

    code = ensure_style(document, "DatumX Code", WD_STYLE_TYPE.PARAGRAPH)
    code.base_style = normal
    code.font.name = "Courier New"
    code.font.size = Pt(9)
    code.font.color.rgb = TEXT
    code.paragraph_format.space_after = Pt(4)
    code.paragraph_format.line_spacing = 1.0


def set_paragraph_spacing(paragraph, before: int = 0, after: int = 0, line: float | None = None) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    if line is not None:
        fmt.line_spacing = line


def add_run(paragraph, text: str, *, font_name: str, size: float, bold: bool = False, italic: bool = False, color: RGBColor = TEXT):
    run = paragraph.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    return run


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color: str = "C6D1C6", size: str = "8") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_paragraph_border(paragraph, edge: str = "bottom", color: str = "C6D1C6", size: str = "10") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = p_bdr.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        p_bdr.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), "1")
    border.set(qn("w:color"), color)


def add_page_field(paragraph) -> None:
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)


def configure_header_footer(section, title: str) -> None:
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    header.clear()
    add_run(header, title.upper(), font_name="Courier New", size=8, bold=True, color=MUTED)
    set_paragraph_spacing(header, after=0)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.clear()
    add_run(footer, f"{title} | Page ", font_name="Courier New", size=8.5, color=MUTED)
    add_page_field(footer)
    set_paragraph_spacing(footer, after=0)


def add_cover_page(document: Document, config: DocumentConfig) -> None:
    if LOGO_PATH.exists():
        logo = document.add_paragraph()
        logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        logo_run = logo.add_run()
        logo_run.add_picture(str(LOGO_PATH), width=Inches(1.2))
        set_paragraph_spacing(logo, before=6, after=8)

    label = document.add_paragraph(style="DatumX Label")
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(label, "DATUMX PROTOCOL DOSSIER", font_name="Courier New", size=9, bold=True, color=PRIMARY)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(title, config.cover_title, font_name="Courier New", size=21, bold=True, color=TEXT)
    set_paragraph_spacing(title, before=18, after=8, line=1.08)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(subtitle, config.cover_subtitle, font_name="Arial", size=12, italic=True, color=MUTED)
    set_paragraph_spacing(subtitle, after=18)

    value = document.add_paragraph()
    value.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(value, config.value_prop, font_name="Arial", size=11.5, color=TEXT)
    set_paragraph_spacing(value, after=14, line=1.18)

    meta = document.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = True
    labels = ["STACK", "LIVE", "PALETTE", "FORMAT"]
    values = [
        config.stack_line,
        config.live_line,
        "Primary #78A37B | Secondary #5F7FA7 | Callout fill #EEF4EE",
        "US Letter | 1 inch margins | Courier New headings | Arial body",
    ]
    for idx, (label_text, value_text) in enumerate(zip(labels, values, strict=True)):
        left = meta.cell(idx, 0)
        right = meta.cell(idx, 1)
        set_cell_shading(left, CALLOUT_FILL)
        set_cell_shading(right, "FFFFFF")
        set_cell_border(left)
        set_cell_border(right)
        left.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        right.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        add_run(left.paragraphs[0], label_text, font_name="Courier New", size=8.5, bold=True, color=PRIMARY)
        add_run(right.paragraphs[0], value_text, font_name="Arial", size=10, color=TEXT)

    note = document.add_paragraph(style="DatumX Meta")
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_run(note, config.intro_note, font_name="Arial", size=10, italic=True, color=MUTED)
    set_paragraph_spacing(note, before=14, after=0)


def parse_markdown(text: str) -> list[tuple[str, object]]:
    lines = [normalize_text(line) for line in text.splitlines()]
    blocks: list[tuple[str, object]] = []
    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("```"):
            fence_lines = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                fence_lines.append(lines[i])
                i += 1
            blocks.append(("code", "\n".join(fence_lines).strip("\n")))
            i += 1
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            blocks.append(("heading", (level, strip_markdown_inline(heading_match.group(2)))))
            i += 1
            continue

        if stripped.startswith(">"):
            quote_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quote_lines.append(strip_markdown_inline(lines[i].strip()[1:].strip()))
                i += 1
            blocks.append(("blockquote", " ".join(quote_lines)))
            continue

        if stripped.startswith("|"):
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                table_lines.append(lines[i].strip())
                i += 1
            rows = []
            for raw in table_lines:
                cells = [strip_markdown_inline(cell) for cell in raw.strip("|").split("|")]
                rows.append(cells)
            if len(rows) >= 2 and all(re.fullmatch(r":?-{3,}:?", cell.replace(" ", "")) for cell in rows[1]):
                rows.pop(1)
            blocks.append(("table", rows))
            continue

        if re.match(r"^(-|\*)\s+", stripped) or re.match(r"^\d+\.\s+", stripped):
            items = []
            ordered = False
            while i < len(lines):
                candidate = lines[i].strip()
                if re.match(r"^(-|\*)\s+", candidate):
                    items.append(strip_markdown_inline(re.sub(r"^(-|\*)\s+", "", candidate)))
                    i += 1
                    continue
                match = re.match(r"^(\d+)\.\s+(.*)$", candidate)
                if match:
                    ordered = True
                    items.append(strip_markdown_inline(match.group(2)))
                    i += 1
                    continue
                break
            blocks.append(("list", (ordered, items)))
            continue

        paragraph_lines = []
        while i < len(lines):
            candidate = lines[i].strip()
            if not candidate:
                break
            if candidate.startswith(("```", ">", "|")) or re.match(r"^(#{1,6})\s+", candidate):
                break
            if re.match(r"^(-|\*)\s+", candidate) or re.match(r"^\d+\.\s+", candidate):
                break
            paragraph_lines.append(strip_markdown_inline(candidate))
            i += 1
        blocks.append(("paragraph", " ".join(paragraph_lines)))
    return blocks


def add_static_contents(document: Document, config: DocumentConfig, headings: list[str]) -> None:
    heading = document.add_paragraph(style="DatumX Section Title")
    add_run(heading, config.static_contents_title, font_name="Courier New", size=14, bold=True, color=TEXT)
    set_paragraph_border(heading, edge="bottom")

    note = document.add_paragraph(style="DatumX Meta")
    add_run(note, "Static contents generated from the Markdown source headings.", font_name="Arial", size=9.5, italic=True, color=MUTED)

    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header_cells = table.rows[0].cells
    set_cell_shading(header_cells[0], TABLE_HEADER_FILL)
    set_cell_shading(header_cells[1], TABLE_HEADER_FILL)
    set_cell_border(header_cells[0])
    set_cell_border(header_cells[1])
    add_run(header_cells[0].paragraphs[0], "#", font_name="Courier New", size=8.5, bold=True, color=PRIMARY)
    add_run(header_cells[1].paragraphs[0], "Section", font_name="Courier New", size=8.5, bold=True, color=PRIMARY)

    for idx, title in enumerate(headings, start=1):
        row = table.add_row().cells
        set_cell_border(row[0])
        set_cell_border(row[1])
        add_run(row[0].paragraphs[0], str(idx), font_name="Courier New", size=8.5, bold=True, color=TEXT)
        add_run(row[1].paragraphs[0], title, font_name="Arial", size=10, color=TEXT)


def add_callout(document: Document, label: str, text: str) -> None:
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    set_cell_border(cell)
    label_paragraph = cell.paragraphs[0]
    add_run(label_paragraph, label.upper(), font_name="Courier New", size=8.5, bold=True, color=PRIMARY)
    body = cell.add_paragraph(style="DatumX Callout")
    add_run(body, text, font_name="Arial", size=10.5, italic=True, color=TEXT)
    set_paragraph_spacing(body, before=1, after=4, line=1.14)


def render_blocks(document: Document, config: DocumentConfig, blocks: list[tuple[str, object]]) -> None:
    first_h1_consumed = False
    for kind, payload in blocks:
        if kind == "heading":
            level, text = payload  # type: ignore[misc]
            if level == 1 and not first_h1_consumed:
                first_h1_consumed = True
                continue
            if level == 2:
                heading = document.add_paragraph(style="DatumX Section Title")
                add_run(heading, text, font_name="Courier New", size=14, bold=True, color=TEXT)
                set_paragraph_border(heading, edge="bottom")
                if text in config.callouts:
                    label, callout_text = config.callouts[text]
                    add_callout(document, label, callout_text)
            elif level == 3:
                paragraph = document.add_paragraph(style="DatumX Subsection")
                add_run(paragraph, text, font_name="Courier New", size=11, bold=True, color=SECONDARY)
            else:
                paragraph = document.add_paragraph(style="DatumX Meta")
                add_run(paragraph, text, font_name="Courier New", size=9.5, bold=True, color=MUTED)
            continue

        if kind == "paragraph":
            paragraph = document.add_paragraph(style="DatumX Body")
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            add_run(paragraph, payload, font_name="Arial", size=10.5, color=TEXT)
            set_paragraph_spacing(paragraph, after=7, line=1.18)
            continue

        if kind == "blockquote":
            add_callout(document, "Companion", str(payload))
            continue

        if kind == "list":
            ordered, items = payload  # type: ignore[misc]
            for idx, item in enumerate(items, start=1):
                paragraph = document.add_paragraph(style="DatumX Body")
                paragraph.paragraph_format.left_indent = Inches(0.22)
                marker = f"{idx}. " if ordered else "- "
                add_run(paragraph, marker, font_name="Courier New", size=9.5, bold=True, color=PRIMARY)
                add_run(paragraph, item, font_name="Arial", size=10.5, color=TEXT)
                set_paragraph_spacing(paragraph, after=4, line=1.14)
            continue

        if kind == "code":
            table = document.add_table(rows=1, cols=1)
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            cell = table.cell(0, 0)
            set_cell_shading(cell, CODE_FILL)
            set_cell_border(cell, color="D7DEE2")
            paragraph = cell.paragraphs[0]
            paragraph.style = document.styles["DatumX Code"]
            add_run(paragraph, str(payload), font_name="Courier New", size=8.8, color=TEXT)
            set_paragraph_spacing(paragraph, before=5, after=5, line=1.0)
            continue

        if kind == "table":
            rows = payload  # type: ignore[assignment]
            if not rows:
                continue
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.alignment = WD_TABLE_ALIGNMENT.LEFT
            table.style = "Table Grid"
            for row_idx, row_values in enumerate(rows):
                for col_idx, value in enumerate(row_values):
                    cell = table.cell(row_idx, col_idx)
                    set_cell_border(cell)
                    if row_idx == 0:
                        set_cell_shading(cell, TABLE_HEADER_FILL)
                    cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
                    paragraph = cell.paragraphs[0]
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    font = "Courier New" if row_idx == 0 else "Arial"
                    size = 8.5 if row_idx == 0 else 9.5
                    add_run(paragraph, value, font_name=font, size=size, bold=row_idx == 0, color=TEXT)
                    set_paragraph_spacing(paragraph, before=2, after=2, line=1.08)


def build_document(config: DocumentConfig) -> None:
    raw_text = config.source_path.read_text(encoding="utf-8", errors="replace")
    blocks = parse_markdown(raw_text)
    headings = [payload[1] for kind, payload in blocks if kind == "heading" and payload[0] == 2]

    document = Document()
    configure_styles(document)
    set_section_margins(document.sections[0])
    configure_header_footer(document.sections[0], config.doc_title)

    add_cover_page(document, config)

    contents_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    set_section_margins(contents_section)
    configure_header_footer(contents_section, config.doc_title)
    add_static_contents(document, config, headings)

    body_section = document.add_section(WD_SECTION_START.NEW_PAGE)
    set_section_margins(body_section)
    configure_header_footer(body_section, config.doc_title)
    render_blocks(document, config, blocks)

    config.output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(config.output_path)


def main() -> None:
    for config in DOCS:
        build_document(config)
        print(f"Generated: {config.output_path}")


if __name__ == "__main__":
    main()
